"""Streamlit UI for exploring workbook formulas and related data."""

from __future__ import annotations

from pathlib import Path
import tempfile
import sys

import streamlit as st
from openpyxl.utils.cell import range_boundaries

if __package__:
    from .dependency_graph import DependencyGraph
    from .formula_parser import CellReference, FormulaParser, FormulaReference
    from .workbook_loader import LoadedWorkbook, WorkbookLoader
else:
    sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
    from formula_map.dependency_graph import DependencyGraph
    from formula_map.formula_parser import CellReference, FormulaParser, FormulaReference
    from formula_map.workbook_loader import LoadedWorkbook, WorkbookLoader


MAX_RELATED_CELLS = 400


def _meaning(formula: str) -> str:
    """Return a short human-readable description for common formulas."""
    upper_formula = formula.upper()
    functions = (
        ("SUM(", "Sum values from the referenced cells"),
        ("AVERAGE(", "Calculate the average of the referenced cells"),
        ("MIN(", "Find the minimum value from the referenced cells"),
        ("MAX(", "Find the maximum value from the referenced cells"),
        ("COUNT(", "Count numeric values from the referenced cells"),
        ("IF(", "Return a value based on a condition"),
    )
    for function, description in functions:
        if function in upper_formula:
            return description
    return "Calculate a value from the referenced cells"


def _reference_values(workbook, reference: FormulaReference):
    worksheet = workbook[reference.sheet]
    min_col, min_row, max_col, max_row = range_boundaries(
        f"{reference.start}:{reference.end}"
    )
    rows = list(
        worksheet.iter_rows(
            min_row=min_row,
            max_row=max_row,
            min_col=min_col,
            max_col=max_col,
            values_only=True,
        )
    )
    headers = [
        worksheet.cell(row=min_row, column=column).column_letter
        for column in range(min_col, max_col + 1)
    ]
    return headers, rows


def _render_reference(workbook, reference: FormulaReference) -> None:
    st.markdown(f"**{reference.sheet}!{reference.start}:{reference.end}**")
    min_col, min_row, max_col, max_row = range_boundaries(
        f"{reference.start}:{reference.end}"
    )
    cell_count = (max_col - min_col + 1) * (max_row - min_row + 1)
    if cell_count > MAX_RELATED_CELLS:
        st.info(
            f"This range contains {cell_count:,} cells. "
            f"Only ranges up to {MAX_RELATED_CELLS} cells are shown."
        )
        return

    headers, rows = _reference_values(workbook, reference)
    st.dataframe(
        {header: [row[index] for row in rows] for index, header in enumerate(headers)},
        use_container_width=True,
        hide_index=True,
    )


def _load_local_workbook(path: Path) -> LoadedWorkbook:
    return WorkbookLoader().load(path)


def _load_uploaded_workbook(uploaded_file) -> LoadedWorkbook:
    suffix = Path(uploaded_file.name).suffix.lower() or ".xlsx"
    with tempfile.NamedTemporaryFile(suffix=suffix, delete=False) as temporary:
        temporary.write(uploaded_file.getvalue())
        path = Path(temporary.name)
    return WorkbookLoader().load(path, keep_vba=suffix in {".xlsm", ".xltm"})


def _workbook_source():
    uploaded_file = st.file_uploader(
        "Upload workbook",
        type=["xlsx", "xlsm", "xltx", "xltm"],
    )
    if uploaded_file is not None:
        return _load_uploaded_workbook(uploaded_file), uploaded_file.name

    local_files = sorted(Path("data/local").glob("*.xlsx"))
    if not local_files:
        st.warning("Upload an Excel workbook to begin.")
        return None, None
    path = st.selectbox("Local workbook", local_files, format_func=lambda item: item.name)
    return _load_local_workbook(path), path.name


def run() -> None:
    st.set_page_config(page_title="FormulaMap", layout="wide")
    st.title("FormulaMap")
    st.caption("Explore Excel formulas, dependencies, and related data.")

    loaded, filename = _workbook_source()
    if loaded is None:
        return

    try:
        workbook = loaded.workbook
        graph = DependencyGraph.from_workbook(loaded)
        parser = FormulaParser()

        st.sidebar.header("Selection")
        sheet_name = st.sidebar.selectbox("Sheet", workbook.sheetnames)
        worksheet = workbook[sheet_name]
        formula_cells = [
            cell.coordinate
            for row in worksheet.iter_rows()
            for cell in row
            if isinstance(cell.value, str) and cell.value.startswith("=")
        ]
        if not formula_cells:
            st.warning(f"No formulas found in {sheet_name}.")
            return

        coordinate = st.sidebar.selectbox("Formula cell", formula_cells)
        target = CellReference(sheet_name, coordinate)
        formula = worksheet[coordinate].value
        references = parser.parse(formula, sheet_name)

        st.subheader(filename)
        st.markdown(f"### {sheet_name}!{coordinate}")
        st.markdown("**Formula**")
        st.code(formula, language="excel")
        st.markdown("**Meaning**")
        st.write(_meaning(formula))

        st.markdown("**Dependencies**")
        dependencies = graph.dependencies_of(target)
        if dependencies:
            st.dataframe(
                {
                    "Sheet": [cell.sheet for cell in dependencies],
                    "Cell": [cell.coordinate for cell in dependencies],
                },
                use_container_width=True,
                hide_index=True,
            )
        else:
            st.info("No cell dependencies found.")

        st.markdown("**Related Data**")
        if references:
            for reference in references:
                _render_reference(workbook, reference)
        else:
            st.info("No rectangular data ranges found in this formula.")
    finally:
        loaded.close()


if __name__ == "__main__":
    run()