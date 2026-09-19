"""Workbook loading with formula-preserving defaults."""

from dataclasses import dataclass
from pathlib import Path

from openpyxl import Workbook, load_workbook


@dataclass
class LoadedWorkbook:
    """An opened workbook and the path it came from."""

    path: Path
    workbook: Workbook

    def close(self) -> None:
        """Release the workbook resources."""
        self.workbook.close()


class WorkbookLoader:
    """Load Excel workbooks without replacing formulas with cached values."""

    SUPPORTED_SUFFIXES = {".xlsx", ".xlsm", ".xltx", ".xltm"}

    def load(self, path: str | Path, *, keep_vba: bool = False) -> LoadedWorkbook:
        workbook_path = Path(path)
        if not workbook_path.is_file():
            raise FileNotFoundError(f"Workbook does not exist: {workbook_path}")
        if workbook_path.suffix.lower() not in self.SUPPORTED_SUFFIXES:
            supported = ", ".join(sorted(self.SUPPORTED_SUFFIXES))
            raise ValueError(
                f"Unsupported workbook type {workbook_path.suffix!r}; expected {supported}"
            )

        workbook = load_workbook(
            workbook_path,
            data_only=False,
            keep_vba=keep_vba,
        )
        return LoadedWorkbook(path=workbook_path, workbook=workbook)